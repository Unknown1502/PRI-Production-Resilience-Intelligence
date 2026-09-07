"""Building deliberately broken workbooks, one defect at a time.

The eighteen error codes each need a file that triggers exactly that error and
nothing else. Eighteen binary ``.xlsx`` fixtures committed to the repository
would be unreviewable — a diff on one says ``Binary files differ`` — and would
go stale the first time a column is added.

So they are built here instead, from a declarative baseline that every test
mutates in one specific way. The defect a test introduces is visible in the
test, next to the assertion about it.
"""

from __future__ import annotations

import copy
from datetime import date, time
from io import BytesIO
from typing import Any

import pytest
from openpyxl import Workbook

from pri.importer.spec import WORKBOOK_SPEC, SheetSpec

__all__ = ["SheetRows", "baseline", "build_xlsx", "workbook_from"]

#: A sheet's rows, each a mapping of column name to value.
SheetRows = list[dict[str, Any]]


def _baseline_rows() -> dict[str, SheetRows]:
    """A minimal, entirely valid production.

    Two shooting days, three scenes, two people, two locations, one piece of
    equipment. Small enough to read in one screen, complete enough that every
    reference resolves and the constraint validator finds nothing.
    """
    return {
        "production": [
            {
                "production_id": "film-test",
                "title": "Test Production",
                "currency": "GBP",
                "timezone": "Europe/London",
                "shoot_start": date(2027, 6, 1),
                "shoot_end": date(2027, 6, 30),
                "reserve_days": "2027-06-20",
                "template_version": "1",
            }
        ],
        "scenes": [
            {
                "scene_id": "S1",
                "number": "1",
                "slug": "INT. ROOM - DAY",
                "description": "One",
                "int_ext": "INT",
                "time_of_day": "DAY",
                "estimated_minutes": 120,
                "location_id": "L1",
                "cast_ids": "P1",
                "crew_ids": "",
                "equipment_ids": "E1",
                "vfx_plate": "FALSE",
                "prerequisite_scene_ids": "",
            },
            {
                "scene_id": "S2",
                "number": "2",
                "slug": "INT. ROOM - NIGHT",
                "description": "Two",
                "int_ext": "INT",
                "time_of_day": "NIGHT",
                "estimated_minutes": 90,
                "location_id": "L1",
                "cast_ids": "P1",
                "crew_ids": "P2",
                "equipment_ids": "E1",
                "vfx_plate": "FALSE",
                "prerequisite_scene_ids": "",
            },
            {
                "scene_id": "S3",
                "number": "3",
                "slug": "EXT. YARD - DAY",
                "description": "Three",
                "int_ext": "EXT",
                "time_of_day": "DAY",
                "estimated_minutes": 60,
                "location_id": "L2",
                "cast_ids": "P1",
                "crew_ids": "",
                "equipment_ids": "",
                "vfx_plate": "FALSE",
                "prerequisite_scene_ids": "S1",
            },
        ],
        "people": [
            {
                "person_id": "P1",
                "name": "Alex Reed",
                "role": "CAST",
                "character": "Alex",
                "daily_rate": 1200,
            },
            {
                "person_id": "P2",
                "name": "Sam Idris",
                "role": "CREW",
                "character": "",
                "daily_rate": 400,
            },
        ],
        "locations": [
            {
                "location_id": "L1",
                "name": "Stage One",
                "kind": "studio",
                "day_rate": 2000,
                "permit_start_time": None,
                "permit_end_time": None,
                "supports_int_ext": "INT",
                "supports_time_of_day": "DAY;NIGHT",
            },
            {
                "location_id": "L2",
                "name": "Back Yard",
                "kind": "exterior",
                "day_rate": 800,
                "permit_start_time": time(7, 0),
                "permit_end_time": time(19, 0),
                "supports_int_ext": "EXT",
                "supports_time_of_day": "DAY",
            },
        ],
        "equipment": [
            {"equipment_id": "E1", "name": "Camera A", "kind": "camera", "daily_rate": 500},
        ],
        "availability": [],
        "schedule": [
            {
                "date": date(2027, 6, 1),
                "call_time": time(8, 0),
                "wrap_time": time(18, 0),
                "wrap_next_day": "FALSE",
                "location_id": "L1",
                "scene_ids": "S1;S2",
                "unit": "MAIN",
            },
            {
                "date": date(2027, 6, 2),
                "call_time": time(9, 0),
                "wrap_time": time(17, 0),
                "wrap_next_day": "FALSE",
                "location_id": "L2",
                "scene_ids": "S3",
                "unit": "MAIN",
            },
        ],
    }


@pytest.fixture()
def baseline() -> dict[str, SheetRows]:
    """A fresh copy of the valid baseline, safe to mutate."""
    return copy.deepcopy(_baseline_rows())


def build_xlsx(
    rows_by_sheet: dict[str, SheetRows],
    *,
    omit_sheets: frozenset[str] = frozenset(),
    omit_columns: dict[str, frozenset[str]] | None = None,
    extra_columns: dict[str, dict[str, Any]] | None = None,
    extra_sheets: dict[str, list[list[Any]]] | None = None,
    hidden_sheets: frozenset[str] = frozenset(),
    hidden_rows: dict[str, frozenset[int]] | None = None,
    merges: dict[str, list[str]] | None = None,
) -> bytes:
    """Write a workbook from plain dicts, with optional deliberate damage.

    Inputs:
        rows_by_sheet:  Sheet name to its data rows.
        omit_sheets:    Sheets to leave out entirely (for E001).
        omit_columns:   Columns to leave out, per sheet (for E002).
        extra_columns:  Unknown columns to add, per sheet (for W006).
        extra_sheets:   Whole sheets PRI does not know, as raw row lists.
        hidden_sheets:  Sheets to mark hidden (for W013).
        hidden_rows:    Excel row numbers to hide, per sheet (for W013).
        merges:         Ranges to merge, per sheet (for W012).

    Outputs:
        The workbook as bytes.
    """
    omit_columns = omit_columns or {}
    extra_columns = extra_columns or {}
    hidden_rows = hidden_rows or {}
    merges = merges or {}

    workbook = Workbook()
    default = workbook.active
    if default is not None:
        workbook.remove(default)

    for sheet_spec in WORKBOOK_SPEC:
        if sheet_spec.name in omit_sheets or sheet_spec.name not in rows_by_sheet:
            continue
        _write(
            workbook,
            sheet_spec,
            rows_by_sheet[sheet_spec.name],
            omitted=omit_columns.get(sheet_spec.name, frozenset()),
            extras=extra_columns.get(sheet_spec.name, {}),
            hidden=sheet_spec.name in hidden_sheets,
            hidden_row_numbers=hidden_rows.get(sheet_spec.name, frozenset()),
            merge_ranges=merges.get(sheet_spec.name, []),
        )

    for name, raw_rows in (extra_sheets or {}).items():
        worksheet = workbook.create_sheet(name)
        for row in raw_rows:
            worksheet.append(row)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write(
    workbook: Workbook,
    sheet_spec: SheetSpec,
    rows: SheetRows,
    *,
    omitted: frozenset[str],
    extras: dict[str, Any],
    hidden: bool,
    hidden_row_numbers: frozenset[int],
    merge_ranges: list[str],
) -> None:
    worksheet = workbook.create_sheet(sheet_spec.name)
    headers = [c.name for c in sheet_spec.columns if c.name not in omitted]
    headers.extend(extras)
    worksheet.append(headers)

    for row in rows:
        merged = {**row, **extras}
        worksheet.append([merged.get(header) for header in headers])

    if hidden:
        worksheet.sheet_state = "hidden"
    for number in hidden_row_numbers:
        worksheet.row_dimensions[number].hidden = True
    for reference in merge_ranges:
        worksheet.merge_cells(reference)


def workbook_from(baseline_rows: dict[str, SheetRows], **damage: Any) -> bytes:
    """Shorthand: ``build_xlsx`` over a baseline that a test has already mutated."""
    return build_xlsx(baseline_rows, **damage)
